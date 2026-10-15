"""Deterministic Word documents (SPECS S10.9, NFR-11).

Every document is stamped with a fixed time (an approval time, never the clock):
core properties ``created`` and ``modified`` and every entry of the .docx
container carry it, and the author is "Blueprint Studio". The same state
therefore always renders to the same bytes.
"""

import io
import zipfile
from collections.abc import Sequence
from datetime import UTC, datetime

from docx import Document
from docx.document import Document as DocumentObject
from docx.shared import Pt

AUTHOR = "Blueprint Studio"


def parse_time(iso: str) -> datetime:
    """An ISO 8601 UTC time from state (``...Z``) as an aware datetime."""
    return datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone(UTC)


def new_document(title: str) -> DocumentObject:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    document.core_properties.title = title
    return document


def table(document: DocumentObject, headers: Sequence[str], rows: Sequence[Sequence[str]]) -> None:
    grid = document.add_table(rows=1, cols=len(headers))
    grid.style = "Table Grid"
    for cell, header in zip(grid.rows[0].cells, headers, strict=True):
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
    for row in rows:
        cells = grid.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = value


def bullets(document: DocumentObject, items: Sequence[str], empty: str = "None.") -> None:
    if not items:
        document.add_paragraph(empty)
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def to_bytes(document: DocumentObject, at: datetime) -> bytes:
    """Save with fixed core properties and a container whose every entry carries ``at``."""
    stamp = at.astimezone(UTC).replace(tzinfo=None, microsecond=0)
    props = document.core_properties
    props.author = AUTHOR
    props.last_modified_by = AUTHOR
    props.created = stamp
    props.modified = stamp
    props.last_printed = stamp
    props.revision = 1
    raw = io.BytesIO()
    document.save(raw)
    return restamp_zip(raw.getvalue(), stamp)


def restamp_zip(data: bytes, at: datetime) -> bytes:
    """The same zip with every entry dated ``at`` and fixed attributes, in the same order."""
    date_time = (at.year, at.month, at.day, at.hour, at.minute, at.second)
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            entry = zipfile.ZipInfo(info.filename, date_time=date_time)
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o600 << 16
            entry.create_system = 0
            target.writestr(entry, source.read(info.filename))
    return out.getvalue()
