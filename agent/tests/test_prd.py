"""T-12: PRD v1.0 — draft, sign-off and first download (S2.10, S2.12, S7.1-S7.4, S10.1-S10.4)."""

import io
import json
import shutil
import subprocess
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from fakes import ScriptedChatModel, prd_out
from langchain_core.messages import AIMessage, HumanMessage
from psycopg_pool import AsyncConnectionPool
from seam import GraphSeam
from service import build_service, client

from blueprint.domain import Design, Gap, NfrTarget, Note, PrdDoc, Requirement
from blueprint.downloads import render
from blueprint.state import SessionState, new_session_state

DESIGN = Design(
    functional_requirements=[
        Requirement(id="FR-1", statement="Users shorten a long URL  (exactly as typed)."),
        Requirement(id="FR-2", statement="Visitors are redirected from a short URL"),
        Requirement(id="FR-3", statement="Owners see their links"),
    ],
    non_functional_requirements=[
        NfrTarget(
            id="NFR-1",
            category="latency",
            statement="Redirects are fast",
            metric="p95",
            target=50,
            unit="ms",
        )
    ],
    out_of_scope=[Note(id="OOS-1", text="Link analytics")],
)
GAPS = [
    Gap(
        id="GAP-1",
        lens="security",
        item_ids=["FR-1"],
        risk="Open creation",
        decision_needed="Who?",
        severity="high",
        status="accepted",
        note="v1",
    ),
    Gap(
        id="GAP-2",
        lens="efficiency",
        item_ids=["NFR-1"],
        risk="No cache",
        decision_needed="Cache?",
        severity="low",
        status="resolved",
        resolved_by=["FR-3"],
    ),
]
GOOD = prd_out(
    {"Short links": ["FR-1", "FR-2"], "Link owners": ["FR-3"]}, {"GAP-1": ["Short links"]}
)


def drafting(*script: Any) -> GraphSeam:
    return GraphSeam(
        ScriptedChatModel.with_script(*script),
        initial={"stage": "prd_draft", "current_step": 5, "confirmed_steps": [1, 2, 3, 4, 5],
                 "design": DESIGN.model_dump(mode="json"), "agent_design": DESIGN.model_dump(mode="json"),
                 "id_counters": {"FR": 3, "NFR": 1, "OOS": 1, "GAP": 2},
                 "gaps": [g.model_dump(mode="json") for g in GAPS]},
    )  # fmt: skip


async def draft(seam: GraphSeam) -> tuple[PrdDoc, AIMessage]:
    await seam.resync()
    state = await seam.state()
    gate = state["messages"][-1]
    assert isinstance(gate, AIMessage)
    return PrdDoc.model_validate(state["prd_draft"]), gate


async def test_the_draft_copies_requirements_verbatim_and_opens_sign_off() -> None:
    seam = drafting(GOOD)
    doc, gate = await draft(seam)

    assert [(r.id, r.statement) for r in doc.functional_requirements] == [
        (r.id, r.statement) for r in DESIGN.functional_requirements
    ]
    assert [n.model_dump() for n in doc.non_functional_requirements] == [
        n.model_dump(exclude={"id"}) | {"id": n.id} for n in DESIGN.non_functional_requirements
    ]
    assert doc.out_of_scope == ["Link analytics"]
    assert [(f.id, f.name, f.requirement_ids) for f in doc.features] == [
        ("F-1", "Short links", ["FR-1", "FR-2"]),
        ("F-2", "Link owners", ["FR-3"]),
    ]
    assert [r.feature_id for r in doc.functional_requirements] == ["F-1", "F-1", "F-2"]
    assert [(g.id, g.feature_ids) for g in doc.open_gaps] == [
        ("GAP-1", ["F-1"])
    ]  # only accepted gaps
    state = await seam.state()
    assert state["stage"] == "prd_signoff"
    assert {g["id"]: g["feature_ids"] for g in state["gaps"]}["GAP-1"] == ["F-1"]
    assert gate.tool_calls[0]["name"] == "sign_off_prd"
    assert gate.tool_calls[0]["args"]["prd"] == state["prd_draft"]
    sent = seam.interview.calls[0].messages[1].text
    assert "Users shorten a long URL  (exactly as typed)." in sent


async def test_an_invalid_draft_is_corrected_once_then_reported() -> None:
    twice = prd_out({"A": ["FR-1", "FR-2"], "B": ["FR-2", "FR-3"]}, {"GAP-1": ["A"]})
    seam = drafting(twice, GOOD)
    doc, _ = await draft(seam)
    assert len(doc.features) == 2
    retry = seam.interview.calls[1].messages[-1]
    assert isinstance(retry, HumanMessage) and "FR-2 is in 2 features" in retry.text

    no_gap = prd_out({"A": ["FR-1", "FR-2", "FR-3"]})
    failing = drafting(no_gap, no_gap)
    await failing.resync()
    state = await failing.state()
    assert state["prd_draft"] is None and state["stage"] == "prd_draft"
    assert "open gap GAP-1 must name at least one feature" in state["last_error"]["message"]


async def test_sign_off_freezes_v1_0_and_rejection_redrafts_with_the_reason() -> None:
    revised = prd_out(
        {"Short links": ["FR-1", "FR-2"], "Owner dashboard": ["FR-3"]}, {"GAP-1": ["Short links"]}
    )
    seam = drafting(GOOD, revised)
    _, gate = await draft(seam)

    await seam.answer(
        str(gate.tool_calls[0]["id"]),
        json.dumps({"approved": False, "reason": "Call the second feature Owner dashboard"}),
    )
    state = await seam.state()
    redraft = PrdDoc.model_validate(state["prd_draft"])
    assert "Call the second feature Owner dashboard" in seam.interview.calls[1].messages[1].text
    assert [(f.id, f.name) for f in redraft.features] == [
        ("F-1", "Short links"),
        ("F-3", "Owner dashboard"),
    ]
    assert state["prd_versions"] == [] and state["stage"] == "prd_signoff"

    second_gate = state["messages"][-1]
    await seam.answer(str(second_gate.tool_calls[0]["id"]), '{"approved": true}')
    state = await seam.state()
    assert state["stage"] == "spec_synthesis"
    [version] = state["prd_versions"]
    assert version["version"] == "1.0" and version["doc"] == state["prd_draft"]
    assert version["approved_at"].endswith("Z")
    frozen = json.dumps(version)

    await seam.send("Can we rename a feature?")  # later turns never touch a signed-off version
    assert json.dumps((await seam.state())["prd_versions"][0]) == frozen


# --- downloads (export seam and REST) ---------------------------------------------------------


def signed_off_state(**extra: Any) -> SessionState:
    out = prd_out(
        {"Short links": ["FR-1", "FR-2"], "Link owners": ["FR-3"]}, {"GAP-1": ["Short links"]}
    )
    from blueprint.prd import PrdDraftOut, build_prd

    doc = build_prd(PrdDraftOut.model_validate(out), DESIGN, [GAPS[0]], {})
    state = new_session_state("s", 1000, True) | {
        "stage": "spec_synthesis",
        "design": DESIGN.model_dump(mode="json"),
        "gaps": [g.model_dump(mode="json") for g in GAPS],
        "prd_draft": doc.model_dump(mode="json"),
        "prd_versions": [
            {
                "version": "1.0",
                "approved_at": "2026-10-14T12:30:00Z",
                "doc": doc.model_dump(mode="json"),
            }
        ],
    }
    return SessionState.model_validate(state | extra)


def headings(data: bytes) -> list[str]:
    return [
        p.text
        for p in Document(io.BytesIO(data)).paragraphs
        if p.style is not None and p.style.name.startswith(("Heading", "Title"))
    ]


def test_prd_v1_0_docx_has_the_s10_4_sections_in_order() -> None:
    data = render(signed_off_state(), "prd-v1.0.docx", "URL shortener")
    assert headings(data) == [
        "URL shortener — PRD v1.0", "Problem Statement", "Goals", "Users and Stakeholders", "Features",
        "F-1: Short links", "F-2: Link owners", "Functional Requirements", "Non-Functional Requirements",
        "Scope", "Out of Scope", "Open Gaps and Accepted Risks", "Diagrams",
    ]  # fmt: skip
    document = Document(io.BytesIO(data))
    assert "approved 2026-10-14 12:30 UTC" in document.paragraphs[1].text
    fr, nfr, gaps = ([[c.text for c in row.cells] for row in t.rows] for t in document.tables)
    assert fr == [["ID", "Feature", "Requirement"], ["FR-1", "F-1: Short links", "Users shorten a long URL  (exactly as typed)."],
                  ["FR-2", "F-1: Short links", "Visitors are redirected from a short URL"], ["FR-3", "F-2: Link owners", "Owners see their links"]]  # fmt: skip
    assert nfr == [
        ["ID", "Category", "Requirement", "Metric", "Target", "Unit"],
        ["NFR-1", "latency", "Redirects are fast", "p95", "50", "ms"],
    ]
    assert gaps == [["ID", "Lens", "Severity", "Affected items", "Risk", "Decision needed", "Note"],
                    ["GAP-1", "security", "high", "FR-1", "Open creation", "Who?", "v1"]]  # fmt: skip
    core = document.core_properties
    assert (core.author, core.created.isoformat(), core.modified.isoformat()) == (
        "Blueprint Studio",
        "2026-10-14T12:30:00+00:00",
        "2026-10-14T12:30:00+00:00",
    )
    assert (
        render(signed_off_state(), "prd-v1.0.docx", "URL shortener") == data
    )  # pure and deterministic


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice is not installed")
def test_prd_docx_opens_in_libreoffice(tmp_path: Path) -> None:
    """NFR-11: LibreOffice loads and lays out the document (converting it to PDF)."""
    source = tmp_path / "PRD v1.0.docx"
    source.write_bytes(render(signed_off_state(), "prd-v1.0.docx", "URL shortener"))
    assert libreoffice_pdf(source, tmp_path).startswith(b"%PDF")


def libreoffice_pdf(source: Path, workdir: Path) -> bytes:
    profile = f"-env:UserInstallation=file://{workdir}/lo-profile"
    command = [
        "soffice",
        profile,
        "--headless",
        "--norestore",
        "--convert-to",
        "pdf",
        "--outdir",
        str(workdir),
        str(source),
    ]
    result = subprocess.run(command, capture_output=True, text=True, timeout=60)
    assert result.returncode == 0, result.stderr
    return source.with_suffix(".pdf").read_bytes()


async def test_downloads_are_listed_and_served_once_available(
    pool: AsyncConnectionPool[Any],
) -> None:
    app = await build_service(pool)
    async with client(app) as http:
        session = (await http.post("/api/sessions", json={"name": "Ünïcode shortener"})).json()
        listing = (await http.get(f"/api/sessions/{session['id']}/downloads")).json()
        assert [i["kind"] for i in listing] == [
            "prd-v1.0.docx", "prd-v1.1.docx", "prd-v1.2.docx", "requirements.docx", "user-stories.docx", "context.md",
            "diagrams.md", "bundle.zip",
        ]  # fmt: skip
        assert not any(i["available"] for i in listing)
        early = await http.get(f"/api/sessions/{session['id']}/downloads/prd-v1.0.docx")
        assert early.status_code == 409 and early.json()["error"]["code"] == "not_available"
        assert (
            await http.get(f"/api/sessions/{session['id']}/downloads/nope.pdf")
        ).status_code == 404

        signed = signed_off_state().model_dump(mode="json")
        await app.state.graph.aupdate_state({"configurable": {"thread_id": session["id"]}}, signed)
        listing = {
            i["kind"]: i
            for i in (await http.get(f"/api/sessions/{session['id']}/downloads")).json()
        }
        assert listing["prd-v1.0.docx"] == {
            "kind": "prd-v1.0.docx",
            "available": True,
            "reason": None,
        }
        assert (
            listing["requirements.docx"]["available"] is False
            and listing["requirements.docx"]["reason"]
        )
        response = await http.get(f"/api/sessions/{session['id']}/downloads/prd-v1.0.docx")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )
    assert (
        response.headers["content-disposition"]
        == "attachment; filename=\"PRD v1.0.docx\"; filename*=UTF-8''PRD%20v1.0.docx"
    )
    assert headings(response.content)[0] == "Ünïcode shortener — PRD v1.0"
