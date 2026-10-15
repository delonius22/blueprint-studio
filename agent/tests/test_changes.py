"""T-14: proposed changes and PRD v1.2 (S2.12, S2.13, S8.5, S8.6, S10.4); Q-11 reading (a)."""

import io
import json
from typing import Any

from docx import Document
from fakes import ScriptedChatModel
from langchain_core.messages import AIMessage
from seam import GraphSeam
from test_prd import signed_off_state
from test_spec import SPEC, story

from blueprint.changes import apply_changes, change_problem
from blueprint.domain import PrdDoc, ProposedChange
from blueprint.downloads import render
from blueprint.prd import current_prd
from blueprint.state import SessionState


def v11_state(changes: list[dict[str, Any]]) -> dict[str, Any]:
    base = signed_off_state()
    v10 = base.prd_versions[0].doc
    v11 = v10.model_copy(
        update={"solution": "A web service", "testing_decisions": ["SEAM-1: HTTP API."]}
    )
    return base.model_dump(mode="json") | {
        "stage": "change_review",
        "prd_draft": v11.model_dump(mode="json"),
        "prd_versions": [
            base.prd_versions[0].model_dump(mode="json"),
            {
                "version": "1.1",
                "approved_at": "2026-10-14T13:00:00Z",
                "doc": v11.model_dump(mode="json"),
            },
        ],
        "proposed_changes": [
            {"id": f"CHG-{i + 1}", "decision": "pending", **c} for i, c in enumerate(changes)
        ],
        "id_counters": {"FR": 3, "NFR": 1, "F": 2, "CHG": len(changes)},
    }


CHANGES = [
    {"section": "goals", "kind": "add", "proposed_text": "Links expire after a year", "reason": "Story US-4 needs it"},
    {"section": "functional_requirements", "target_id": "FR-2", "kind": "modify",
     "current_text": "Visitors are redirected from a short URL", "proposed_text": "Visitors are redirected with HTTP 301", "reason": "Precision"},
    {"section": "scope", "kind": "remove", "current_text": "Creating and resolving short links", "reason": "Duplicated"},
]  # fmt: skip


async def reviewing(changes: list[dict[str, Any]]) -> tuple[GraphSeam, str]:
    seam = GraphSeam(ScriptedChatModel.with_script(), initial=v11_state(changes))
    await seam.resync()
    gate = (await seam.state())["messages"][-1]
    assert isinstance(gate, AIMessage) and gate.tool_calls[0]["name"] == "review_changes"
    return seam, str(gate.tool_calls[0]["id"])


async def test_review_changes_lists_every_pending_change() -> None:
    seam, _ = await reviewing(CHANGES)
    gate = (await seam.state())["messages"][-1]
    assert [c["id"] for c in gate.tool_calls[0]["args"]["changes"]] == ["CHG-1", "CHG-2", "CHG-3"]
    await seam.resync()  # a page load does not open the card twice
    assert (
        sum(
            bool(isinstance(m, AIMessage) and m.tool_calls)
            for m in (await seam.state())["messages"]
        )
        == 1
    )


async def test_approved_changes_produce_v1_2_and_rejected_ones_are_dropped() -> None:
    seam, card = await reviewing(CHANGES)
    before = json.dumps((await seam.state())["prd_versions"])

    decisions = [{"change_id": "CHG-1", "decision": "approved"}, {"change_id": "CHG-2", "decision": "approved"},
                 {"change_id": "CHG-3", "decision": "rejected", "reason": "Keep it"}]  # fmt: skip
    await seam.answer(card, json.dumps(decisions))

    state = await seam.state()
    assert state["stage"] == "ticket_breakdown"
    v10, v11, v12 = state["prd_versions"]
    assert json.dumps([v10, v11]) == json.dumps(json.loads(before))  # earlier versions unchanged
    assert v12["version"] == "1.2"
    doc = PrdDoc.model_validate(v12["doc"])
    assert doc.goals == ["Short links that always redirect", "Links expire after a year"]
    assert [r.statement for r in doc.functional_requirements if r.id == "FR-2"] == [
        "Visitors are redirected with HTTP 301"
    ]
    assert doc.scope == ["Creating and resolving short links"]  # the rejected removal was dropped
    assert doc.solution == "A web service" and doc.testing_decisions == ["SEAM-1: HTTP API."]
    assert [c["decision"] for c in state["proposed_changes"]] == [
        "approved",
        "approved",
        "rejected",
    ]
    # S8.6: every later stage and download reads the newest version
    assert current_prd(state["prd_versions"]) == doc


async def test_no_v1_2_when_every_change_is_rejected() -> None:
    seam, card = await reviewing(CHANGES[:1])
    await seam.answer(
        card, json.dumps([{"change_id": "CHG-1", "decision": "rejected", "reason": "No"}])
    )
    state = await seam.state()
    assert [v["version"] for v in state["prd_versions"]] == ["1.0", "1.1"]
    assert state["stage"] == "ticket_breakdown"
    assert current_prd(state["prd_versions"]) == PrdDoc.model_validate(
        state["prd_versions"][1]["doc"]
    )


def test_adding_and_removing_requirements_issues_ids_and_keeps_features_whole() -> None:
    doc = signed_off_state().prd_versions[0].doc
    add = ProposedChange(id="CHG-1", section="functional_requirements", target_id="F-2", kind="add",
                         proposed_text="Owners delete links", reason="r")  # fmt: skip
    remove = ProposedChange(
        id="CHG-2", section="functional_requirements", target_id="FR-3", kind="remove", reason="r"
    )
    counters = {"FR": 3}
    changed = apply_changes(doc, [add, remove], counters)
    assert [(r.id, r.feature_id) for r in changed.functional_requirements] == [
        ("FR-1", "F-1"),
        ("FR-2", "F-1"),
        ("FR-4", "F-2"),
    ]
    assert [(f.id, f.requirement_ids) for f in changed.features] == [
        ("F-1", ["FR-1", "FR-2"]),
        ("F-2", ["FR-4"]),
    ]
    assert counters == {"FR": 4}
    assert doc.functional_requirements[2].id == "FR-3"  # the source is not modified


def test_changes_the_rules_cannot_apply_are_rejected_at_synthesis() -> None:
    doc = signed_off_state().prd_versions[0].doc

    def problem(**fields: Any) -> str | None:
        return change_problem(ProposedChange(id="CHG-1", reason="r", **fields), doc)

    assert (
        problem(section="open_gaps", kind="remove", target_id="GAP-1")
        == "open gaps come from the design review and cannot be changed"
    )
    assert "must group at least one requirement" in str(
        problem(section="features", kind="add", proposed_text="New")
    )
    assert "needs target_id naming its feature" in str(
        problem(section="functional_requirements", kind="add", proposed_text="x")
    )
    assert "must equal an entry of goals" in str(
        problem(section="goals", kind="remove", current_text="Not there")
    )
    assert problem(section="users", kind="add", proposed_text="Admins") is None


async def test_an_inapplicable_proposal_is_corrected_at_synthesis() -> None:
    from test_spec import SEAMS, synthesizing

    bad = SPEC | {"proposed_changes": [{"section": "open_gaps", "target_id": "GAP-1", "kind": "remove", "reason": "Fixed"}],
                  "user_stories": [story("F-1", ["FR-1", "FR-2"]), story("F-2", ["FR-3"])]}  # fmt: skip
    seam = synthesizing(bad, SPEC, SEAMS)
    await seam.send("go")
    retry = seam.interview.calls[1].messages[-1]
    assert (
        "proposed change 1 cannot be applied: open gaps come from the design review" in retry.text
    )


def test_prd_v1_2_docx_lists_the_approved_changes() -> None:
    state = v11_state(CHANGES)
    doc = apply_changes(
        PrdDoc.model_validate(state["prd_versions"][1]["doc"]),
        [ProposedChange.model_validate(state["proposed_changes"][0])],
        {},
    )
    state["prd_versions"].append(
        {
            "version": "1.2",
            "approved_at": "2026-10-14T14:00:00Z",
            "doc": doc.model_dump(mode="json"),
        }
    )
    state["proposed_changes"][0]["decision"] = "approved"
    state["proposed_changes"][1]["decision"] = "rejected"
    document = Document(
        io.BytesIO(render(SessionState.model_validate(state), "prd-v1.2.docx", "URL shortener"))
    )
    headings = [
        p.text for p in document.paragraphs if p.style is not None and p.style.name == "Heading 1"
    ]
    assert headings[-1] == "Changes from v1.1"
    rows = [[c.text for c in row.cells] for row in document.tables[-1].rows]
    assert rows == [["ID", "Section", "Target", "Change", "Current text", "Proposed text", "Reason"],
                    ["CHG-1", "goals", "", "add", "", "Links expire after a year", "Story US-4 needs it"]]  # fmt: skip


async def test_the_ticket_planner_reads_v1_2_as_the_current_prd() -> None:
    from fixtures import complete_state
    from test_tickets import ticket

    done = complete_state()
    keep = ("design", "gaps", "prd_versions", "prd_draft")
    initial = {k: done.model_dump(mode="json")[k] for k in keep} | {
        "stage": "change_review",
        "current_step": 5,
        "tickets": [],
        "tickets_approved_at": None,
        "proposed_changes": [{"id": "CHG-1", "decision": "pending", **CHANGES[0]}],
        "id_counters": {"FR": 3, "NFR": 1, "F": 2, "US": 3, "CHG": 1},
    }
    tickets = {"tickets": [ticket("a", ["US-1"]), ticket("b", ["US-2", "US-3"], ["a"])]}
    seam = GraphSeam(ScriptedChatModel.with_script(tickets), initial=initial)
    await seam.resync()
    card = (await seam.state())["messages"][-1].tool_calls[0]["id"]

    await seam.answer(str(card), json.dumps([{"change_id": "CHG-1", "decision": "approved"}]))

    state = await seam.state()
    assert [v["version"] for v in state["prd_versions"]] == ["1.0", "1.1", "1.2"]
    planner_input = seam.interview.calls[0].messages[1].text
    assert "Links expire after a year" in planner_input  # only v1.2 has this goal
    assert state["stage"] == "ticket_gate"
