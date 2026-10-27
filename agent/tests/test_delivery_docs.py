"""T-16: Requirements.docx and UserStories.docx (S10.3, S10.5, S10.6, NFR-11)."""

import io
import shutil
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from fixtures import TICKETS, complete_state
from psycopg_pool import AsyncConnectionPool
from service import build_service, client
from test_prd import libreoffice_pdf

from blueprint.exports.delivery_docx import build_waves
from blueprint.exports.downloads import render


def lines(data: bytes) -> list[str]:
    return [p.text for p in Document(io.BytesIO(data)).paragraphs]


def tables(data: bytes) -> list[list[list[str]]]:
    return [
        [[c.text for c in row.cells] for row in t.rows] for t in Document(io.BytesIO(data)).tables
    ]


def test_requirements_docx_lists_requirements_with_coverage_and_the_gap_appendix() -> None:
    data = render(complete_state(), "requirements.docx", "URL shortener")
    assert [line for line in lines(data) if line in ("Functional Requirements", "Non-Functional Requirements", "Appendix: Open Gaps and Accepted Risks")] == [
        "Functional Requirements", "Non-Functional Requirements", "Appendix: Open Gaps and Accepted Risks",
    ]  # fmt: skip
    fr, nfr, gaps = tables(data)
    assert fr == [
        ["ID", "Feature", "Requirement", "Covering stories"],
        ["FR-1", "F-1: Short links", "Users shorten a long URL  (exactly as typed).", "US-1"],
        ["FR-2", "F-1: Short links", "Visitors are redirected from a short URL", "US-2"],
        ["FR-3", "F-2: Link owners", "Owners see their links", "US-3"],
    ]
    assert nfr == [
        ["ID", "Category", "Requirement", "Covering stories"],
        ["NFR-1", "latency", "Redirects are fast", "US-1"],
    ]
    assert gaps[1] == ["GAP-1", "security", "high", "F-1", "FR-1", "Open creation", "Who?", "v1"]


def test_user_stories_docx_has_stories_by_feature_tickets_and_build_waves() -> None:
    data = render(complete_state(), "user-stories.docx", "URL shortener")
    text = lines(data)
    assert text.index("F-1: Short links") < text.index(
        "US-1: As an anonymous user, I want to shorten a long URL, so that I can share it"
    )
    assert text.index(
        "US-3: As a link owner, I want to see my links, so that I can manage them"
    ) > text.index("F-2: Link owners")
    assert "Requirements: FR-1, NFR-1. Delivered by: T-01." in text
    assert (
        text.index("Tickets")
        < text.index("T-01: Walking skeleton")
        < text.index("T-02: Redirect")
        < text.index("T-03: Owner list")
    )
    assert "Blocked by: T-01" in text and "Blocked by: None (can start immediately)" in text
    assert "☐ GET /{code} answers 301" in text and "Delivers: US-2" in text
    assert tables(data)[-1] == [
        ["Wave", "Tickets", "Blocked by"],
        ["1", "T-01", "—"],
        ["2", "T-02, T-03", "T-01"],
    ]


def test_build_waves_come_from_the_blocking_edges() -> None:
    assert [[t.id for t in w] for w in build_waves(TICKETS)] == [["T-01"], ["T-02", "T-03"]]


def test_both_documents_are_deterministic() -> None:
    for kind in ("requirements.docx", "user-stories.docx"):
        assert render(complete_state(), kind, "URL shortener") == render(
            complete_state(), kind, "URL shortener"
        )


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice is not installed")
@pytest.mark.parametrize("kind", ["requirements.docx", "user-stories.docx"])
def test_both_open_in_libreoffice(tmp_path: Path, kind: str) -> None:
    source = tmp_path / kind
    source.write_bytes(render(complete_state(), kind, "URL shortener"))
    assert libreoffice_pdf(source, tmp_path).startswith(b"%PDF")


async def test_both_are_409_until_the_session_is_complete(pool: AsyncConnectionPool[Any]) -> None:
    app = await build_service(pool)
    async with client(app) as http:
        session = (await http.post("/api/sessions", json={"name": "URL shortener"})).json()
        url = f"/api/sessions/{session['id']}/downloads"
        config = {"configurable": {"thread_id": session["id"]}}
        await app.state.graph.aupdate_state(
            config,
            complete_state(stage="ticket_gate", tickets_approved_at=None).model_dump(mode="json"),
        )
        for kind in ("requirements.docx", "user-stories.docx"):
            response = await http.get(f"{url}/{kind}")
            assert (
                response.status_code == 409
                and "ticket breakdown is approved" in response.json()["error"]["message"]
            )

        await app.state.graph.aupdate_state(config, complete_state().model_dump(mode="json"))
        names = {}
        for kind in ("requirements.docx", "user-stories.docx"):
            response = await http.get(f"{url}/{kind}")
            assert response.status_code == 200
            names[kind] = response.headers["content-disposition"]
    assert 'filename="Requirements.docx"' in names["requirements.docx"]
    assert 'filename="UserStories.docx"' in names["user-stories.docx"]
