"""T-13: PRD v1.1 with to-spec, frozen with the seams it proposes (S2.11, S8.1-S8.4, S10.4, S12.3)."""

import io
import json
from typing import Any

from docx import Document
from fakes import ScriptedChatModel
from langchain_core.messages import AIMessage, HumanMessage, SystemMessage
from seam import GraphSeam
from test_prd import DESIGN, GAPS, signed_off_state

from blueprint.domain import SIGNED_OFF_SECTIONS, PrdDoc
from blueprint.downloads import render
from blueprint.exports.prd_docx import ATTRIBUTION
from blueprint.state import SessionState


def story(feature: str, requirements: list[str], want: str = "to shorten a URL") -> dict[str, Any]:
    return {
        "feature_id": feature,
        "actor": "an anonymous user",
        "feature": want,
        "benefit": "I can share it",
        "acceptance_criteria": ["A short code is returned"],
        "requirement_ids": requirements,
    }


SPEC = {
    "solution": "A small web service that issues short codes and redirects.",
    "user_stories": [
        story("F-1", ["FR-1", "NFR-1"]),
        story("F-1", ["FR-2"], "to be redirected"),
        story("F-2", ["FR-3"], "to see my links"),
    ],
    "implementation_decisions": ["Codes are base62"],
    "further_notes": ["Analytics may come later"],
    "proposed_changes": [],
}
SEAMS = {
    "seams": [
        {
            "name": "HTTP API",
            "description": "Drive the public API end to end",
            "covers": ["US-1", "US-2", "US-3"],
        }
    ]
}


def synthesizing(*script: Any, **extra: Any) -> GraphSeam:
    state = signed_off_state().model_dump(mode="json")
    seam = GraphSeam(ScriptedChatModel.with_script(*script), initial={k: state[k] for k in (
        "stage", "design", "gaps", "prd_draft", "prd_versions", "current_step", "confirmed_steps",
    )} | {"id_counters": {"FR": 3, "NFR": 1, "OOS": 1, "GAP": 2, "F": 2}, "current_step": 5} | extra)  # fmt: skip
    return seam


async def test_the_synthesizer_runs_to_spec_verbatim_without_touching_signed_off_sections() -> None:
    seam = synthesizing(SPEC, SEAMS)
    await seam.send("We agreed on everything; please write the spec.")

    system, human = seam.interview.calls[0].messages
    assert isinstance(system, SystemMessage) and seam.assets.to_spec in system.text
    assert (
        system.text.index("Vendored verbatim from mattpocock-skills 1.2.3") > 0
    )  # with its attribution
    assert "do NOT ask the user any questions" in system.text
    assert isinstance(human, HumanMessage)
    assert '"problem_statement": "Long URLs are hard to share."' in human.text  # v1.0
    assert "Users shorten a long URL  (exactly as typed)." in human.text  # design record
    assert "User: We agreed on everything; please write the spec." in human.text  # transcript

    state = await seam.state()
    v10 = state["prd_versions"][0]["doc"]
    draft = state["prd_draft"]
    for section in SIGNED_OFF_SECTIONS:
        assert json.dumps(draft[section]) == json.dumps(v10[section]), section
    assert [(s["id"], s["feature_id"]) for s in draft["user_stories"]] == [
        ("US-1", "F-1"),
        ("US-2", "F-1"),
        ("US-3", "F-2"),
    ]
    assert draft["solution"] == SPEC["solution"]
    assert draft["further_notes"] == [
        "Analytics may come later",
        "GAP-1: Open creation Decision needed: Who? Note: v1",
    ]
    assert [v["version"] for v in state["prd_versions"]] == ["1.0", "1.1"]
    assert state["stage"] == "ticket_breakdown"  # no approval: the tickets are next


async def test_stories_and_changes_are_validated_with_one_retry() -> None:
    bad = SPEC | {
        "user_stories": [story("F-9", ["FR-1"]), story("F-1", ["FR-7"])],
        "proposed_changes": [{"section": "solution", "kind": "modify", "reason": "x"}],
    }
    seam = synthesizing(bad, SPEC, SEAMS)
    await seam.send("go")
    retry = seam.interview.calls[1].messages[-1]
    assert isinstance(retry, HumanMessage)
    for problem in ("story 1 names feature F-9, which does not exist", "story 2 names FR-7, which is not an FR or NFR",
                    "FR-2 is covered by no user story", "proposed change 1 names solution, which is not a signed-off section"):  # fmt: skip
        assert problem in retry.text
    assert (await seam.state())["stage"] == "ticket_breakdown"

    failing = synthesizing(bad, bad)
    await failing.send("go")
    state = await failing.state()
    assert state["stage"] == "spec_synthesis" and state["last_error"]["code"] == "invalid_output"
    assert state["prd_draft"] == state["prd_versions"][0]["doc"]  # nothing half-written


async def test_the_proposed_seams_freeze_v1_1_without_an_approval_card() -> None:
    seam = synthesizing(SPEC, SEAMS)
    await seam.send("go")

    state = await seam.state()
    assert state["stage"] == "ticket_breakdown"  # no proposed changes: on to the tickets
    v10, v11 = state["prd_versions"]
    assert v11["version"] == "1.1"
    assert v11["doc"]["testing_decisions"] == [
        "SEAM-1: HTTP API. Drive the public API end to end Covers US-1, US-2, US-3."
    ]
    for section in SIGNED_OFF_SECTIONS:
        assert json.dumps(v11["doc"][section]) == json.dumps(v10["doc"][section])
    assert [s["id"] for s in state["seams"]] == ["SEAM-1"]
    assert not any(isinstance(m, AIMessage) and m.tool_calls for m in state["messages"])
    assert any(
        "PRD v1.1 is frozen with 1 test seam (HTTP API)" in m.text for m in state["messages"]
    )


async def test_proposed_changes_are_applied_after_v1_1_as_v1_2() -> None:
    change = {
        "section": "goals",
        "kind": "add",
        "proposed_text": "Links expire",
        "reason": "Stories need it",
    }
    seam = synthesizing(SPEC | {"proposed_changes": [change]}, SEAMS)
    await seam.send("go")
    state = await seam.state()
    assert [(c["id"], c["section"], c["decision"]) for c in state["proposed_changes"]] == [
        ("CHG-1", "goals", "approved")
    ]
    assert [v["version"] for v in state["prd_versions"]] == ["1.0", "1.1", "1.2"]
    assert "Links expire" in state["prd_versions"][-1]["doc"]["goals"]
    assert state["stage"] == "ticket_breakdown"


def test_prd_v1_1_docx_adds_the_spec_sections_and_ends_with_the_attribution() -> None:
    base = signed_off_state()
    v10 = base.prd_versions[0]
    doc = PrdDoc.model_validate(v10.doc.model_dump() | {
        "solution": "A small web service.",
        "user_stories": [{"id": "US-1", **story("F-1", ["FR-1"])}],
        "implementation_decisions": ["Codes are base62"],
        "testing_decisions": ["SEAM-1: HTTP API."],
        "further_notes": ["GAP-1: Open creation Decision needed: Who? Note: v1"],
    })  # fmt: skip
    state = SessionState.model_validate(base.model_dump() | {"prd_versions": [
        v10.model_dump(), {"version": "1.1", "approved_at": "2026-10-14T13:00:00Z", "doc": doc.model_dump()}]})  # fmt: skip
    document = Document(io.BytesIO(render(state, "prd-v1.1.docx", "URL shortener")))
    headings = [
        p.text
        for p in document.paragraphs
        if p.style is not None and p.style.name.startswith("Heading 1")
    ]
    assert headings[-5:] == [
        "Solution",
        "User Stories",
        "Implementation Decisions",
        "Testing Decisions",
        "Further Notes",
    ]
    assert document.paragraphs[-1].text == ATTRIBUTION
    assert (
        "US-1 (F-1; FR-1): As an anonymous user, I want to shorten a URL, so that I can share it"
        in [p.text for p in document.paragraphs]
    )
    assert DESIGN.functional_requirements[0].statement in [
        c.text for row in document.tables[0].rows for c in row.cells
    ]
    assert GAPS[0].id in [c.text for row in document.tables[2].rows for c in row.cells]
