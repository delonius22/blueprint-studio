"""The design as Mermaid diagrams, in the downloads (S11.3; architecture, data model, sequences).

agent/tests/expected/diagrams/ is shared with web/tests/diagram-parity.test.ts: the browser
and the downloads must draw the same source, and the web test checks it is valid Mermaid.
"""

import io
import json

from conftest import REPO_ROOT
from docx import Document
from fixtures import complete_state
from test_prd import signed_off_state

from blueprint.domain.models import Design
from blueprint.exports import diagrams
from blueprint.exports.downloads import render, unmet

EXPECTED = REPO_ROOT / "agent" / "tests" / "expected" / "diagrams"
DESIGN = Design.model_validate(json.loads((EXPECTED / "design.json").read_text()))


def test_the_architecture_matches_the_shared_expected_source() -> None:
    assert diagrams.architecture(DESIGN) == (EXPECTED / "architecture.mmd").read_text()


def test_the_data_model_matches_the_shared_expected_source() -> None:
    assert diagrams.data_model(DESIGN) == (EXPECTED / "data-model.mmd").read_text()


def test_each_flow_has_a_sequence_matching_the_shared_expected_source() -> None:
    for flow in DESIGN.flows:
        expected = (EXPECTED / f"sequence-{flow.id}.mmd").read_text()
        assert diagrams.sequence(flow, DESIGN.components) == expected, flow.id


def test_only_what_was_recorded_is_drawn() -> None:
    assert diagrams.all_diagrams(Design()) == []
    titles = [title for title, _ in diagrams.all_diagrams(DESIGN)]
    assert titles == [
        "Architecture",
        "Data model",
        "Sequence: FLOW-1 Resolve a code (read)",
        "Sequence: FLOW-2 Shorten a URL (write)",
    ]
    entities_only = Design(entities=DESIGN.entities)
    assert [t for t, _ in diagrams.all_diagrams(entities_only)] == ["Data model"]


def test_diagrams_md_is_available_as_soon_as_there_is_something_to_draw() -> None:
    interviewing = signed_off_state().model_copy(update={"stage": "interview"})
    assert (
        unmet(interviewing, "diagrams.md") == "available once entities or components are recorded"
    )
    drawn = interviewing.model_copy(
        update={"design": interviewing.design.model_copy(update={"components": DESIGN.components})}
    )
    assert unmet(drawn, "diagrams.md") is None


def test_diagrams_md_has_a_mermaid_block_per_diagram() -> None:
    text = render(complete_state(), "diagrams.md", "URL shortener").decode("utf-8")
    assert text.startswith("# URL shortener — Diagrams\n")
    assert text.count("```mermaid\n") == 3
    for heading in (
        "## Architecture",
        "## Data model",
        "## Sequence: FLOW-1 Follow a short link (read)",
    ):
        assert f"\n{heading}\n" in text
    design = Design.model_validate(complete_state().design.model_dump())
    assert diagrams.data_model(design).rstrip("\n") in text
    assert text.endswith("```\n")


def test_the_prd_carries_the_mermaid_source_of_every_diagram() -> None:
    document = Document(io.BytesIO(render(complete_state(), "prd-v1.1.docx", "URL shortener")))
    paragraphs = [p.text for p in document.paragraphs]
    start = paragraphs.index("Diagrams")
    assert paragraphs[start + 2 : start + 4] == [
        "Architecture",
        diagrams.architecture(complete_state().design).rstrip("\n"),
    ]
    sources = [p for p in document.paragraphs if p.runs and p.runs[0].font.name == "Courier New"]
    assert len(sources) == 3
    assert paragraphs.index("Solution") > start  # before the v1.1 sections
