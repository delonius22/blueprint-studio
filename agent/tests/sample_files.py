"""Attachment files built in code, so the tests hold no binary fixtures."""

import base64
import io
import urllib.parse
import zlib

import docx

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
)

BRIEF_MD = """# Link shortener

## Requirements
Users shorten a long URL.
Visitors are redirected from a short URL.
Redirects must be fast.

## Out of scope
Link analytics dashboards.

```mermaid
flowchart LR
  CMP_1(["Browser"]) --> CMP_2["API"]
```
"""


def make_docx(blocks: list[tuple[str, str]], table: list[list[str]] | None = None) -> bytes:
    """A .docx of (style, text) paragraphs, e.g. ("Heading 1", "Goals"), then an optional table."""
    document = docx.Document()
    for style, text in blocks:
        document.add_paragraph(text, style=style)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(grid.rows, table, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def make_pdf(pages: list[str]) -> bytes:
    """A minimal PDF with one line of text per page ("" for a page with no text layer)."""
    objects: list[bytes] = []
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font = 3 + 2 * len(pages)
    for i, text in enumerate(pages):
        content = f"BT /F1 12 Tf 72 712 Td ({text}) Tj ET".encode() if text else b""
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R "
            f"/Resources << /Font << /F1 {font} 0 R >> >> >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    out.writelines(b"%010d 00000 n \n" % offset for offset in offsets)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    )
    return out.getvalue()


DRAWIO_MODEL = """<mxGraphModel><root>
  <mxCell id="0"/><mxCell id="1" parent="0"/>
  <mxCell id="a" value="Web &lt;b&gt;client&lt;/b&gt;" style="shape=actor;" vertex="1" parent="1"/>
  <mxCell id="b" value="Orders DB" style="shape=cylinder3;whiteSpace=wrap;" vertex="1" parent="1"/>
  <mxCell id="e" style="edgeStyle=orthogonal;dashed=1;" edge="1" source="a" target="b" parent="1"/>
  <mxCell id="l" value="place order" style="edgeLabel;" vertex="1" parent="e"/>
</root></mxGraphModel>"""


def make_drawio(compressed: bool = False) -> bytes:
    if not compressed:
        return f'<mxfile><diagram name="Backend">{DRAWIO_MODEL}</diagram></mxfile>'.encode()
    deflater = zlib.compressobj(wbits=-15)
    packed = deflater.compress(urllib.parse.quote(DRAWIO_MODEL).encode()) + deflater.flush()
    text = base64.b64encode(packed).decode()
    return f'<mxfile><diagram name="Backend">{text}</diagram></mxfile>'.encode()
